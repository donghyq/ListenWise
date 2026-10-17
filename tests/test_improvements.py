"""Regression coverage for profiles, source boundaries, audio and exports."""
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch, Mock
from pathlib import Path
from xml.etree import ElementTree as ET

from tests import test_api as fixtures


class ImprovementsTest(unittest.TestCase):
    setUpClass = classmethod(fixtures.LawFlowApiTest.setUpClass.__func__)
    tearDownClass = classmethod(fixtures.LawFlowApiTest.tearDownClass.__func__)

    def test_output_floor_is_optional_and_mode_aware(self):
        self.assertEqual(self.main.output_floor(1000, 'condense', 'none', 0.3), 0)
        self.assertEqual(self.main.output_floor(1000, 'condense', 'auto', 0.3), 300)
        self.assertEqual(self.main.output_floor(1000, 'adapt', 'auto', 0.3), 1000)
        self.assertEqual(self.main.output_floor(1000, 'adapt', 'custom', 0.42), 420)

    def test_project_keeps_research_and_output_preferences(self):
        project = self.client.post('/api/projects', json={
            'name': '可配置策略',
            'minimum_output_mode': 'none',
            'minimum_output_ratio': 0.4,
            'web_research_mode': 'augment',
        })
        self.assertEqual(project.status_code, 201)
        data = self.client.get('/api/projects/' + project.json()['id']).json()['project']
        self.assertEqual(data['minimum_output_mode'], 'none')
        self.assertEqual(data['minimum_output_ratio'], 0.4)
        self.assertEqual(data['web_research_mode'], 'augment')

    def make_content(self):
        project = self.client.post('/api/projects', json={'name': '回归', 'scenario': 'daily_brief', 'transform_mode': 'condense', 'verification_mode': 'source_only'}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title': '素材', 'content': '这是原始材料，需要保留事实和限制条件。企业应当按照实际情况评估。'}).json()
        document = self.client.get(f"/api/documents/{source['id']}").json()
        ids = [b['id'] for b in document['blocks'] if b['kind'] == 'paragraph']
        content = self.client.post(f"/api/projects/{project['id']}/skill-host-contents", json={'document_id':source['id'], 'title':'测试讲稿', 'markdown':'好的，以下是为你整理的讲稿：\n\n# 讲稿\n\n这是**材料**正文。请参考[说明](https://example.com)。\n\n如果需要调整，请告诉我。', 'source_block_ids':ids}).json()
        return project, source, ids, content

    def test_provider_error_message_explains_openai_quota(self):
        message = self.main.provider_error_message('{"error":{"code":"credit_balance_exhausted","message":"You have no credits remaining"}}')
        self.assertIn('额度不足', message)
        self.assertNotIn('credit_balance_exhausted', message)

    def test_version_endpoint_reports_public_release_update(self):
        import httpx
        response = httpx.Response(
            200,
            json={
                'tag_name': 'v0.2.0',
                'html_url': 'https://github.com/donghyq/ListenWise/releases/tag/v0.2.0',
            },
            request=httpx.Request('GET', 'https://api.github.com/repos/donghyq/ListenWise/releases/latest'),
        )
        with patch.object(self.main.httpx, 'get', return_value=response):
            result = self.client.get('/api/app/version')
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['current_version'], '0.1.0')
        self.assertEqual(result.json()['latest_version'], '0.2.0')
        self.assertTrue(result.json()['update_available'])
        self.assertTrue(result.json()['check_succeeded'])
        self.assertEqual(self.client.get('/api/health').json()['version'], '0.1.0')

    def test_version_check_failure_never_blocks_local_app(self):
        import httpx
        with patch.object(self.main.httpx, 'get', side_effect=httpx.ConnectError('offline')):
            result = self.client.get('/api/app/version')
        self.assertEqual(result.status_code, 200)
        self.assertFalse(result.json()['update_available'])
        self.assertFalse(result.json()['check_succeeded'])

    def test_existing_database_is_backed_up_before_first_schema_migration(self):
        import sqlite3
        legacy_dir = self.temp_dir / 'legacy-upgrade'
        legacy_dir.mkdir()
        legacy_db = legacy_dir / 'app.db'
        with sqlite3.connect(legacy_db) as conn:
            conn.execute("CREATE TABLE legacy_marker (value TEXT)")
            conn.execute("INSERT INTO legacy_marker VALUES ('keep-me')")
        with patch.object(self.main, 'DATA_DIR', legacy_dir), patch.object(self.main, 'DB_PATH', legacy_db):
            self.main.init_db()
            backups = list((legacy_dir / 'backups').glob('app-before-schema-*.db'))
        self.assertEqual(len(backups), 1)
        with sqlite3.connect(backups[0]) as conn:
            self.assertEqual(conn.execute('SELECT value FROM legacy_marker').fetchone()[0], 'keep-me')

    def test_learning_progress_migration_backs_up_existing_v1_database(self):
        old_dir = self.temp_dir / 'progress-upgrade'
        old_dir.mkdir()
        database = old_dir / 'app.db'
        with sqlite3.connect(database) as conn:
            conn.execute('CREATE TABLE app_metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL, updated_at TEXT NOT NULL)')
            conn.execute("INSERT INTO app_metadata VALUES ('schema_version', '1', 'previous')")
            conn.execute('CREATE TABLE existing_content (value TEXT)')
            conn.execute("INSERT INTO existing_content VALUES ('keep-me')")
        with patch.object(self.main, 'DATA_DIR', old_dir), patch.object(self.main, 'DB_PATH', database):
            self.main.init_db()
        backups = list((old_dir / 'backups').glob('app-before-schema-2-*.db'))
        self.assertEqual(len(backups), 1)
        with sqlite3.connect(backups[0]) as conn:
            self.assertEqual(conn.execute("SELECT value FROM app_metadata WHERE key = 'schema_version'").fetchone()[0], '1')
        with sqlite3.connect(database) as conn:
            self.assertEqual(conn.execute('SELECT value FROM existing_content').fetchone()[0], 'keep-me')
            self.assertEqual(conn.execute("SELECT value FROM app_metadata WHERE key = 'schema_version'").fetchone()[0], '2')
            self.assertIsNotNone(conn.execute("SELECT name FROM sqlite_master WHERE name = 'learning_progress'").fetchone())

    def test_chatgpt_mcp_exposes_core_lawflow_tools(self):
        import asyncio
        tools = asyncio.run(self.main.chatgpt_mcp.list_tools())
        names = {tool.name for tool in tools}
        self.assertTrue({
            'list_lawflow_projects', 'add_lawflow_text_source', 'get_lawflow_source_context',
            'confirm_lawflow_structure', 'update_lawflow_verification_task',
            'save_lawflow_chatgpt_draft', 'create_lawflow_daily_brief', 'export_lawflow_project',
        }.issubset(names))

    def test_chatgpt_mcp_route_requires_deployment_token(self):
        old = os.environ.pop('LAWFLOW_MCP_TOKEN', None)
        try:
            response = self.client.post('/mcp/', json={'jsonrpc':'2.0','id':1,'method':'initialize','params':{}})
            self.assertEqual(response.status_code, 503)
        finally:
            if old is not None:
                os.environ['LAWFLOW_MCP_TOKEN'] = old

    def test_chatgpt_app_handoff_only_returns_selected_source_blocks(self):
        project = self.client.post('/api/projects', json={'name':'Plus 协作','scenario':'daily_brief','transform_mode':'condense','verification_mode':'source_only'}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title':'素材','content':'监管部门发布了新的数据合规提示，企业需要关注适用范围和实施时间。'}).json()
        document = self.client.get('/api/documents/' + source['id']).json()
        source_ids = [block['id'] for block in document['blocks'] if block['kind'] == 'paragraph']
        result = self.client.post(f"/api/projects/{project['id']}/chatgpt-handoff", json={'document_id':source['id'],'title':'五分钟速听','source_block_ids':source_ids})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['source_block_ids'], source_ids)
        self.assertIn('唯一材料', result.json()['prompt'])
        self.assertNotIn('API Key', result.json()['prompt'])

    def test_desktop_launcher_uses_an_available_port(self):
        from scripts import lawflow_desktop
        import socket
        original = lawflow_desktop.os.environ.get('LAWFLOW_PORT')
        lawflow_desktop.os.environ['LAWFLOW_PORT'] = '18081'
        try:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
                listener.bind(('127.0.0.1', 18081))
                listener.listen(1)
                with self.assertRaises(RuntimeError):
                    lawflow_desktop.choose_port()
        finally:
            if original is None:
                lawflow_desktop.os.environ.pop('LAWFLOW_PORT', None)
            else:
                lawflow_desktop.os.environ['LAWFLOW_PORT'] = original

    def test_desktop_launcher_migrates_legacy_projects_only_into_empty_data_dir(self):
        from scripts import lawflow_desktop
        legacy = self.temp_dir / 'legacy-data'
        target = self.temp_dir / 'desktop-data'
        legacy.mkdir()
        target.mkdir()
        with sqlite3.connect(legacy / 'app.db') as conn:
            conn.execute('CREATE TABLE projects (id TEXT)')
            conn.execute("INSERT INTO projects VALUES ('legacy-project')")
        (legacy / 'projects').mkdir()
        (legacy / 'projects' / 'note.txt').write_text('legacy', encoding='utf-8')
        self.assertTrue(lawflow_desktop.migrate_legacy_data(target, legacy))
        self.assertTrue((target / 'projects' / 'note.txt').is_file())
        self.assertFalse(lawflow_desktop.migrate_legacy_data(target, legacy))

    def test_custom_profile_persistence_and_prompt(self):
        data = {'name':'自定义', 'description':'问答式', 'instruction':'先提出具体问题，再逐步解释，保持克制的表达。'}
        response = self.client.post('/api/narrative/profiles', json=data)
        self.assertEqual(response.status_code, 201)
        pid = response.json()['id']
        self.main.init_db()  # Existing-database migration must retain profiles.
        self.assertEqual(self.main.get_style_profiles()[pid]['instruction'], data['instruction'])
        self.assertEqual(self.client.put('/api/narrative/profiles/law_podcast_v4', json=data).status_code, 400)
        data['name'] = '修改名称'
        self.assertEqual(self.client.put('/api/narrative/profiles/' + pid, json=data).status_code, 200)
        request = self.main.NarrativeOutlineRequest(document_id='d', title='测试', source_block_ids=['b'], style_profile=pid, transform_mode='condense', target_duration=3)
        raw = {'sections':[{'heading':'问题', 'source_block_ids':['b']}], 'opening_angle':'先提出问题'}
        blocks = [{'id':'b', 'kind':'paragraph', 'text':'材料正文', 'source_locator':'段落 1'}]
        with patch.object(self.main, 'model_chat', return_value=json.dumps(raw)) as model:
            outline = self.main.create_narrative_outline(request, blocks)
            self.assertIn(data['instruction'], model.call_args.args[0][-1]['content'])
        self.assertEqual(len(outline['sections']), 1)
        self.assertLessEqual(outline['target_total_words'], 720)

    def test_profile_markdown_export_and_import(self):
        created = self.client.post('/api/narrative/profiles', json={
            'name': '可迁移画像',
            'description': '适合专题学习',
            'instruction': '先说明具体问题，再按背景、规则和实务场景展开，保留必要的事实边界。',
        }).json()
        exported = self.client.get('/api/narrative/profiles/' + created['id'] + '/export')
        self.assertEqual(exported.status_code, 200)
        self.assertIn('# 可迁移画像', exported.text)
        imported = self.client.post('/api/narrative/profiles/import', json={'markdown': exported.text})
        self.assertEqual(imported.status_code, 201)
        self.assertEqual(imported.json()['name'], '可迁移画像')
        self.assertIn('具体问题', imported.json()['instruction'])

    def test_podcast_docx_template_sets_layout(self):
        output = self.temp_dir / 'podcast.docx'
        self.main.markdown_to_docx('# 标题\n\n正文段落。', output, template='podcast')
        from docx import Document
        document = Document(output)
        paragraph = document.paragraphs[1]
        self.assertEqual(paragraph.paragraph_format.line_spacing, 3.0)
        self.assertEqual(paragraph.paragraph_format.first_line_indent.pt, 22)
        self.assertEqual(document.sections[0].left_margin.inches, 1.25)

    def test_regenerate_section_preserves_other_source_mappings(self):
        project = self.client.post('/api/projects', json={'name': '章节重做', 'scenario': 'topic_learning'}).json()
        source = self.client.post('/api/projects/' + project['id'] + '/text-sources', json={'title': '素材', 'content': '第一节材料内容，需要保留背景、事实和规则边界。第二节材料内容，需要说明实务场景、责任分工和后续核验事项。'}).json()
        document = self.client.get('/api/documents/' + source['id']).json()
        ids = [block['id'] for block in document['blocks'] if block['kind'] == 'paragraph']
        outline = {'title': '章节稿', 'opening_angle': '', 'closing_angle': '', 'sections': [
            {'id': 'section-1', 'heading': '第一节', 'purpose': '', 'key_points': [], 'source_block_ids': [ids[0]], 'target_words': 100},
            {'id': 'section-2', 'heading': '第二节', 'purpose': '', 'key_points': [], 'source_block_ids': [ids[-1]], 'target_words': 100},
        ]}
        with self.main.db() as conn:
            outline_id = 'outline-regenerate'
            content_id = 'content-regenerate'
            timestamp = self.main.now_iso()
            conn.execute("INSERT INTO narrative_outlines VALUES (?, ?, ?, '', ?, '法律从业者', 'law_podcast_v4', 'short', ?, ?, 'confirmed', ?, ?)", (outline_id, project['id'], source['id'], '章节稿', json.dumps(ids), json.dumps(outline, ensure_ascii=False), timestamp, timestamp))
            markdown = '# 章节稿\n\n## 第一节\n\n旧第一节。\n\n## 第二节\n\n旧第二节。'
            mappings = [{'section_id': 'section-1', 'heading': '第一节', 'source_block_ids': [ids[0]]}, {'section_id': 'section-2', 'heading': '第二节', 'source_block_ids': [ids[-1]]}]
            conn.execute("INSERT INTO narrative_contents VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'pending_review', '', ?, ?)", (content_id, outline_id, project['id'], source['id'], '章节稿', markdown, json.dumps(mappings), json.dumps({}), timestamp, timestamp))
        with patch.object(self.main, 'model_chat', return_value='重做后的第一节。'):
            response = self.client.post('/api/narrative-contents/' + content_id + '/sections/section-1/regenerate')
        self.assertEqual(response.status_code, 200)
        state = self.client.get('/api/projects/' + project['id']).json()['narrative_contents'][0]
        self.assertEqual({item['section_id'] for item in state['section_sources']}, {'section-1', 'section-2'})
        self.assertIn('重做后的第一节', state['markdown'])

    def test_outline_prompt_uses_actual_source_ids_not_legacy_placeholder(self):
        request = self.main.NarrativeOutlineRequest(document_id='d', title='测试', source_block_ids=['actual:paragraph'], transform_mode='condense', target_duration=3)
        blocks = [{'id':'actual:paragraph', 'kind':'paragraph', 'text':'材料正文', 'source_locator':'段落 1'}]
        raw = {'sections':[{'heading':'问题','source_block_ids':['actual:paragraph']}]}
        with patch.object(self.main, 'model_chat', return_value=json.dumps(raw)) as model:
            self.main.create_narrative_outline(request, blocks)
        prompt = model.call_args.args[0][-1]['content']
        self.assertIn('actual:paragraph', prompt)
        self.assertNotIn('b-00001', prompt)

    def test_confirmed_chapter_structure_controls_writing_and_selected_public_source(self):
        project = self.client.post('/api/projects', json={'name': '长文学习', 'scenario': 'topic_learning', 'verification_mode': 'source_only'}).json()
        primary = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title': '原始材料', 'content': '原始材料说明背景和关键事实，同时写明适用条件与事实边界。'}).json()
        public = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title': '公开背景', 'content': '公开资料补充背景知识，并说明相关概念的含义和通常适用场景。'}).json()
        primary_ids = [block['id'] for block in self.client.get('/api/documents/' + primary['id']).json()['blocks']]
        public_ids = [block['id'] for block in self.client.get('/api/documents/' + public['id']).json()['blocks'] if block['kind'] == 'paragraph']
        with self.main.db() as conn:
            conn.execute('UPDATE source_documents SET source_url = ? WHERE id = ?', ('https://example.com/public', public['id']))
        plan = self.client.post(f"/api/projects/{project['id']}/plans", json={'source_document_id': primary['id'], 'include_audio': False}).json()
        chapter = {**plan['chapters'][0], 'title': '我确认的学习章节'}
        self.assertEqual(self.client.put(f"/api/plans/{plan['id']}/confirm", json={'chapters': [chapter]}).status_code, 200)
        request_body = {'document_id': primary['id'], 'source_plan_id': plan['id'], 'title': '专题学习',
                        'source_block_ids': primary_ids + public_ids, 'supplemental_document_ids': [public['id']],
                        'target_words': 1200, 'transform_mode': 'enrich', 'web_research_mode': 'augment'}
        with patch.object(self.main, 'model_chat') as model:
            response = self.client.post(f"/api/projects/{project['id']}/narrative-outlines", json=request_body)
        self.assertEqual(response.status_code, 201, response.text)
        model.assert_not_called()
        outline = response.json()['outline']
        self.assertEqual(outline['sections'][0]['heading'], '我确认的学习章节')
        self.assertEqual(outline['target_total_words'], 1200)
        self.assertEqual(sum(section['target_words'] for section in outline['sections']), 1200)
        self.assertEqual(outline['transform_mode'], 'enrich')
        self.assertTrue(set(public_ids).issubset(outline['sections'][0]['source_block_ids']))
        self.assertEqual(self.client.put(f"/api/narrative-outlines/{response.json()['id']}/confirm", json={'outline': outline}).status_code, 200)
        with patch.object(self.main, 'model_chat', return_value='测试生成内容。') as model:
            generated = self.client.post(f"/api/narrative-outlines/{response.json()['id']}/contents")
        self.assertEqual(generated.status_code, 201, generated.text)
        self.assertIn('公开资料补充背景知识', model.call_args.args[0][-1]['content'])
        self.assertIn('我确认的学习章节', generated.json()['markdown'])
        self.assertEqual(self.client.post(f"/api/projects/{project['id']}/narrative-outlines", json={**request_body, 'web_research_mode': 'off'}).status_code, 400)

    def test_flat_article_outline_uses_distinct_source_sentences_and_respects_length(self):
        request = self.main.NarrativeOutlineRequest(document_id='doc', title='原文主题', source_block_ids=['h', 'p1', 'p2', 'p3', 'p4'], target_words=1000)
        blocks = [{'id': 'h', 'document_id': 'doc', 'kind': 'heading', 'text': '原文主题'},
                  *[{'id': f'p{i}', 'document_id': 'doc', 'kind': 'paragraph', 'text': f'第{i}个不同的事实说明了这篇文章在实际场景中的核心情况和限制条件。'} for i in range(1, 5)]]
        chapter = {'title': '第01章 · 原文主题', 'question': '这篇文章说明了什么？', 'source_block_ids': [block['id'] for block in blocks], 'enabled': True}
        with patch.object(self.main, 'source_dossier', return_value='素材'):
            outline = self.main.outline_from_confirmed_plan(request, [chapter], blocks)
        section = outline['sections'][0]
        self.assertEqual(outline['target_total_words'], 1000)
        self.assertEqual(section['target_words'], 1000)
        self.assertGreaterEqual(len(section['key_points']), 2)
        self.assertNotIn('原文主题', section['key_points'])
        self.assertTrue(all('事实说明' in point for point in section['key_points']))

    def test_web_search_reports_provider_failure_separately_from_no_match(self):
        import httpx
        with patch.object(self.main, '_search_ddg', side_effect=httpx.ConnectTimeout('timeout')), \
             patch.object(self.main, '_search_bing_rss', return_value=[{'title': '当的汉语解释', 'summary': '', 'url': 'https://example.com/dictionary'}]):
            result = self.client.post('/api/web-search', json={'query': '当中国AI企业走上安理会讲台'})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['status'], 'limited')
        self.assertEqual(result.json()['results'], [])
        self.assertIn('检索服务', result.json()['notice'])

    def test_heading_selection_uses_full_section_and_rejects_overlap(self):
        blocks = [
            {'id': 'h1', 'sequence_no': 1, 'kind': 'heading', 'heading_level': 1, 'text': '第一部分'},
            {'id': 'p1', 'sequence_no': 2, 'kind': 'paragraph', 'text': '前一部分'},
            {'id': 'h2', 'sequence_no': 3, 'kind': 'heading', 'heading_level': 1, 'text': '第二部分'},
            {'id': 'h3', 'sequence_no': 4, 'kind': 'heading', 'heading_level': 2, 'text': '第二部分的细节'},
            {'id': 'p2', 'sequence_no': 5, 'kind': 'paragraph', 'text': '后一部分细节'},
        ]
        chapters = self.main.create_plan_chapters(blocks, 'lexcast', ['h2'])
        self.assertEqual(chapters[0]['source_block_ids'], ['h2', 'h3', 'p2'])
        with self.assertRaises(self.main.HTTPException):
            self.main.create_plan_chapters(blocks, 'lexcast', ['h2', 'h3'])

    def test_plan_respects_optional_audio_choice(self):
        project = self.client.post('/api/projects', json={'name': '音频可选', 'audio_enabled': True}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={
            'title': '学习素材', 'content': '这份学习素材说明了背景、适用条件和主要问题，适合生成一节学习内容。',
        }).json()
        base = {'source_document_id': source['id']}
        text_only = self.client.post(f"/api/projects/{project['id']}/plans", json={**base, 'include_audio': False})
        with_audio = self.client.post(f"/api/projects/{project['id']}/plans", json={**base, 'include_audio': True})
        self.assertEqual(text_only.status_code, 201)
        self.assertEqual(with_audio.status_code, 201)
        self.assertFalse(text_only.json()['include_audio'])
        self.assertTrue(with_audio.json()['include_audio'])

    def test_learning_progress_survives_reopening_and_tracks_current_sections(self):
        project, _source, _ids, content = self.make_content()
        content_id = content['id']
        markdown = '# 学习内容\n\n## 第一节\n第一节正文。\n\n## 第二节\n第二节正文。'
        saved = self.client.put(f'/api/narrative-contents/{content_id}', json={'markdown': markdown})
        self.assertEqual(saved.status_code, 200)
        endpoint = f'/api/narrative-contents/{content_id}/progress'
        self.assertEqual(self.client.put(endpoint, json={'section_id': 'section-1', 'completed': True}).status_code, 200)
        moved = self.client.put(endpoint, json={'section_id': 'section-2'})
        self.assertEqual(moved.status_code, 200)
        self.assertEqual(self.client.put(endpoint, json={'section_id': 'section-3'}).status_code, 400)
        reopened = self.client.get(f"/api/projects/{project['id']}").json()
        progress = next(item for item in reopened['narrative_contents'] if item['id'] == content_id)['reading_progress']
        self.assertEqual(progress['last_section_id'], 'section-2')
        self.assertEqual(progress['completed_section_ids'], ['section-1'])
        self.client.put(f'/api/narrative-contents/{content_id}', json={'markdown': '# 学习内容\n\n## 第一节\n保留的正文。'})
        reopened = self.client.get(f"/api/projects/{project['id']}").json()
        progress = next(item for item in reopened['narrative_contents'] if item['id'] == content_id)['reading_progress']
        self.assertEqual(progress['last_section_id'], '')
        self.assertEqual(progress['completed_section_ids'], ['section-1'])

    def test_profile_extraction_is_draft_and_validates_response(self):
        before = len(self.client.get('/api/narrative/profiles').json())
        draft = {'name':'叙事', 'description':'自然', 'instruction':'用具体问题切入，解释术语并用自然的承接句组织段落。'}
        with patch.object(self.main, 'model_chat', return_value=json.dumps(draft)):
            result = self.client.post('/api/narrative/profile-draft', json={'transcript':'播客素材。' * 40})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['status'], 'draft')
        self.assertEqual(len(self.client.get('/api/narrative/profiles').json()), before)
        with patch.object(self.main, 'model_chat', return_value='{"name":"empty"}'):
            result = self.client.post('/api/narrative/profile-draft', json={'transcript':'素材' * 60})
        self.assertEqual(result.status_code, 502)

    def test_transcription_requires_consent_and_configuration(self):
        with patch.object(self.main, 'get_internal_provider_settings', return_value={}):
            response = self.client.post('/api/narrative/transcribe', files={'file':('clip.mp3', b'audio', 'audio/mpeg')})
        self.assertEqual(response.status_code, 400)
        self.assertIn('允许', response.json()['detail'])
        config = {'allow_source_upload':True, 'asr_base_url':'https://asr.example/v1', 'asr_model':'asr', 'asr_api_key':'secret'}
        import httpx
        from unittest.mock import AsyncMock
        response = httpx.Response(200, json={'text':'转写结果'}, request=httpx.Request('POST', 'https://asr.example'))
        with patch.dict('os.environ', {}, clear=True), patch.object(self.main, 'get_internal_provider_settings', return_value=config), patch.object(httpx.AsyncClient, 'post', new=AsyncMock(return_value=response)):
            result = self.client.post('/api/narrative/transcribe', files={'file':('clip.mp3', b'audio', 'audio/mpeg')})
        self.assertEqual(result.json()['transcript'], '转写结果')

    def test_source_validation_and_budget_no_silent_truncation(self):
        self.assertEqual(self.main.read_blocks('missing', []), [])
        request = self.main.NarrativeOutlineRequest(document_id='d', title='测试', source_block_ids=['b'], transform_mode='condense')
        blocks = [{'id':'b', 'kind':'paragraph', 'text':'正文' * 100, 'source_locator':'1'}]
        with self.assertRaises(self.main.HTTPException):
            self.main.normalize_narrative_outline({'sections':[{'heading':'内容','source_block_ids':['invented']}]}, request, blocks)
        with self.assertRaises(self.main.HTTPException):
            self.main.source_dossier(blocks, max_total_chars=50)
        self.assertIn('正文' * 100, self.main.source_dossier(blocks, max_chars_per_block=10))

    def test_confirm_rejects_invalid_outline_and_writing_hints_not_published(self):
        project, source, ids, content = self.make_content()
        row = self.client.get('/api/projects/' + project['id']).json()['narrative_outlines'][0]
        invalid = {'sections':[{'heading':'不合法','source_block_ids':['foreign']}]}
        self.assertEqual(self.client.put('/api/narrative-outlines/' + row['id'] + '/confirm', json={'outline':invalid}).status_code, 400)
        outline = {'title':'标题', 'opening_angle':'请从问题切入', 'closing_angle':'请总结', 'sections':[{'id':'s1','heading':'正文','source_block_ids':ids,'target_words':300}]}
        with patch.object(self.main, 'model_chat', return_value='实际生成的正文。') as model:
            text, _, _ = self.main.generate_narrative_markdown(outline, self.main.read_blocks(source['id']), '读者', 'law_podcast_v4')
        self.assertNotIn('请从问题切入', text)
        self.assertIn('请从问题切入', model.call_args.args[0][-1]['content'])

    def test_podcast_profile_strips_written_markers_and_rejects_truncated_condensation(self):
        project, source, ids, content = self.make_content()
        # 各章 target_words < 200：跳过逐章压缩，让整稿压缩兜底被触发。
        outline = {'title': '标题', 'opening_angle': '', 'closing_angle': '', 'target_total_words': 600,
                   'sections': [{'id': 's1', 'heading': '第一章', 'source_block_ids': ids, 'target_words': 100},
                                {'id': 's2', 'heading': '第二章', 'source_block_ids': ids, 'target_words': 100}]}
        blocks = self.main.read_blocks(source['id'])
        body1 = '**巴西。** 2026年8月12日，卢拉总统签署四项能源法令。' + '第一章展开的完整口播正文。' * 150
        body2 = '第二章的口播正文，用连贯段落讲清楚。' * 100
        # 压缩输出被 token 截断（丢掉第二章）时必须拒绝采纳。
        truncated_condensed = '# 标题\n\n## 第一章\n\n' + '压缩后只剩第一章的长正文。' * 200
        outputs = ['欢迎收听本期节目，今天讲两件事。', body1 + '结尾收束。', body2 + '结尾收束。', truncated_condensed]
        with patch.object(self.main, 'model_chat', side_effect=outputs):
            text, _, _ = self.main.generate_narrative_markdown(outline, blocks, '读者', 'law_podcast_v4')
        self.assertNotIn('**', text)
        # 截断的整稿压缩被守卫拒绝，原稿两章都保留。
        self.assertIn('第一章展开的完整口播正文', text)
        self.assertIn('第二章的口播正文', text)

    def test_podcast_profile_generates_opening_before_first_chapter(self):
        project, source, ids, content = self.make_content()
        outline = {'title': '标题', 'opening_angle': '', 'closing_angle': '', 'target_total_words': 600,
                   'sections': [{'id': 's1', 'heading': '第一章', 'source_block_ids': ids, 'target_words': 300},
                                {'id': 's2', 'heading': '第二章', 'source_block_ids': ids, 'target_words': 300}]}
        blocks = self.main.read_blocks(source['id'])
        body = '第一章口播正文，讲清楚关键事实。' * 30
        model = Mock(side_effect=['欢迎收听本期节目，今天我们聊两件事。', body, '第二章口播正文，承接上文讲清楚。'])
        with patch.object(self.main, 'model_chat', model):
            text, sources, _ = self.main.generate_narrative_markdown(outline, blocks, '读者', 'law_podcast_v4')
        self.assertIn('## 开场', text)
        self.assertIn('欢迎收听本期节目', text)
        self.assertLess(text.index('## 开场'), text.index('## 第一章'))
        self.assertEqual(sources[0]['heading'], '开场')
        # 开场白提示词只做引入预告，不展开事实；第一次调用是开场白，第二次才是章节
        opening_prompt = model.call_args_list[0][0][0][-1]['content']
        self.assertIn('开场白', opening_prompt)
        self.assertIn('第一章；第二章', opening_prompt)

    def test_podcast_single_section_skips_separate_opening(self):
        project, source, ids, content = self.make_content()
        outline = {'title': '标题', 'opening_angle': '', 'closing_angle': '', 'target_total_words': 500,
                   'sections': [{'id': 's1', 'heading': '第一章', 'source_block_ids': ids, 'target_words': 300}]}
        model = Mock(return_value='单章正文，直接讲透。')
        with patch.object(self.main, 'model_chat', model):
            text, _, _ = self.main.generate_narrative_markdown(outline, self.main.read_blocks(source['id']), '读者', 'law_podcast_v4')
        self.assertNotIn('## 开场', text)
        self.assertEqual(model.call_count, 1)

    def test_podcast_profile_no_length_limit_keeps_full_section(self):
        project, source, ids, content = self.make_content()
        outline = {'title': '标题', 'opening_angle': '', 'closing_angle': '', 'target_total_words': 500,
                   'sections': [{'id': 's1', 'heading': '第一章', 'source_block_ids': ids, 'target_words': 300}]}
        blocks = self.main.read_blocks(source['id'])
        # 材料扎实的长章（远超 target_words）：podcast-rewriter 观感下不压缩、保留完整深度。
        long_section = '先看巴西的能源法令，这是本章的关键事实，深度展开不能删。' * 120
        model = Mock(side_effect=[long_section])
        with patch.object(self.main, 'model_chat', model):
            text, _, _ = self.main.generate_narrative_markdown(outline, blocks, '读者', 'law_podcast_v4')
        # 长章完整保留，没有触发任何压缩调用（单章大纲不单独生成开场白）。
        self.assertIn(long_section[:30], text)
        self.assertIn(long_section[-30:], text)
        self.assertEqual(model.call_count, 1)
        # 提示词中不出现字数目标区间。
        prompt = model.call_args.args[0][-1]['content']
        self.assertNotIn('目标区间', prompt)
        self.assertIn('不设上限', prompt)

    def test_audio_edit_preview_and_narrative_only_export(self):
        project, source, ids, content = self.make_content()
        result = self.client.post(f"/api/projects/{project['id']}/audio-scripts", json={'narrative_content_id':content['id']})
        self.assertEqual(result.status_code, 201)
        audio = result.json()
        self.assertNotIn('**', audio['script'])
        self.assertNotIn('https://', audio['script'])
        with patch.object(self.main, 'speech_chunk', return_value=(b'preview', 'test')):
            response = self.client.post('/api/audio-outputs/' + audio['id'] + '/synthesize', json={'preview':True})
        self.assertEqual(response.content, b'preview')
        self.assertFalse(self.client.get('/api/audio-outputs/' + audio['id']).json()['audio_available'])
        old_path = self.main.DATA_DIR / 'old.mp3'; old_path.write_bytes(b'old')
        with self.main.db() as conn:
            conn.execute("UPDATE audio_outputs SET audio_path=?, status='ready' WHERE id=?", (str(old_path), audio['id']))
        conn.close()
        self.assertEqual(self.client.put('/api/audio-outputs/' + audio['id'], json={'script':'修改后的稿件。'}).status_code, 200)
        self.assertFalse(old_path.exists())
        updated = self.client.get('/api/audio-outputs/' + audio['id']).json()
        self.assertFalse(updated['audio_available'])
        self.client.put('/api/narrative-contents/' + content['id'], json={'markdown':'更新后的正文', 'status':'needs_revision'})
        self.assertEqual(self.client.get('/api/audio-outputs/' + audio['id']).json()['status'], 'source_changed')
        result = self.client.post(f"/api/projects/{project['id']}/exports")
        self.assertEqual(result.status_code, 201)
        output = Path(result.json()['output_dir'])
        self.assertTrue((output / '讲稿-01.docx').is_file())
        self.assertTrue((output / '口播脚本-01.txt').is_file())
        evidence = json.loads((output / 'evidence-map.json').read_text())
        self.assertTrue(evidence[0]['section_sources'])

    def test_generation_drops_model_supplied_section_heading(self):
        outline = {'title': '测试讲稿', 'opening_angle': '', 'closing_angle': '', 'sections': [
            {'id': 's1', 'heading': '第一节 适用范围', 'purpose': '', 'key_points': [], 'source_block_ids': ['b'], 'target_words': 100},
        ]}
        blocks = [{'id': 'b', 'kind': 'paragraph', 'text': '材料正文', 'source_locator': '段落 1'}]
        with patch.object(self.main, 'model_chat', return_value='## 第一节 适用范围\n\n正文内容。'):
            markdown, _, _ = self.main.generate_narrative_markdown(outline, blocks, '读者', 'law_podcast_v4')
        self.assertEqual(markdown.count('第一节 适用范围'), 1)
        self.assertIn('正文内容', markdown)

    def test_duplicate_section_headings_are_collapsed_for_display_and_export(self):
        project, source, ids, content = self.make_content()
        with self.main.db() as conn:
            conn.execute(
                'UPDATE narrative_contents SET markdown=? WHERE id=?',
                ('# 测试讲稿\n\n## 第一节\n\n## 第一节\n\n第一节正文。\n\n## 第二节\n\n第二节正文。', content['id']),
            )
        conn.close()
        state = self.client.get('/api/projects/' + project['id']).json()
        markdown = state['narrative_contents'][0]['markdown']
        self.assertEqual(markdown.count('## 第一节'), 1)
        self.assertEqual(markdown.count('## 第二节'), 1)

        exported = self.client.post('/api/narrative-contents/' + content['id'] + '/export')
        exported_markdown = Path(exported.json()['markdown_path']).read_text(encoding='utf-8')
        self.assertTrue(exported_markdown.startswith('待审核｜仅供内部审阅'))
        self.assertEqual(exported_markdown.count('## 第一节'), 1)
        from docx import Document
        exported_docx = Document(exported.json()['docx_path'])
        self.assertIn('待审核', exported_docx.paragraphs[0].text)

    def test_audio_script_can_target_a_single_section(self):
        project, source, ids, content = self.make_content()
        with self.main.db() as conn:
            conn.execute(
                'UPDATE narrative_contents SET markdown=? WHERE id=?',
                ('# 测试讲稿\n\n## 第一节 适用范围\n\n第一节正文，说明规则的适用对象。\n\n## 第二节 备案要求\n\n第二节正文，说明备案与评估手续。', content['id']),
            )
        conn.close()

        whole = self.client.post(f"/api/projects/{project['id']}/audio-scripts", json={'narrative_content_id': content['id']})
        self.assertEqual(whole.status_code, 201)
        self.assertIn('第一节正文', whole.json()['script'])
        self.assertIn('第二节正文', whole.json()['script'])
        self.assertEqual(whole.json()['section_heading'], '')

        section = self.client.post(
            f"/api/projects/{project['id']}/audio-scripts",
            json={'narrative_content_id': content['id'], 'section_heading': '第二节 备案要求'},
        )
        self.assertEqual(section.status_code, 201)
        self.assertIn('第二节正文', section.json()['script'])
        self.assertNotIn('第一节正文', section.json()['script'])
        self.assertIn('第二节 备案要求', section.json()['title'])
        self.assertIn('第二节 备案要求', section.json()['script'])

        missing = self.client.post(
            f"/api/projects/{project['id']}/audio-scripts",
            json={'narrative_content_id': content['id'], 'section_heading': '不存在的章节'},
        )
        self.assertEqual(missing.status_code, 404)

    def test_audio_script_falls_back_when_naturalization_model_is_unavailable(self):
        project, _, _, content = self.make_content()
        result = self.client.post(
            f"/api/projects/{project['id']}/audio-scripts",
            json={"narrative_content_id": content["id"], "naturalize": True},
        )
        self.assertEqual(result.status_code, 201)
        self.assertTrue(result.json()["naturalization_skipped"])
        self.assertFalse(result.json()["naturalized"])

    def test_project_deletion_removes_managed_audio_file(self):
        project, _, _, content = self.make_content()
        audio = self.client.post(
            f"/api/projects/{project['id']}/audio-scripts",
            json={'narrative_content_id': content['id']},
        ).json()
        audio_dir = self.main.DATA_DIR / 'audio'
        audio_dir.mkdir(parents=True, exist_ok=True)
        audio_path = audio_dir / 'managed-output.mp3'
        audio_path.write_bytes(b'audio')
        with self.main.db() as conn:
            conn.execute(
                "UPDATE audio_outputs SET audio_path=?, status='ready' WHERE id=?",
                (str(audio_path), audio['id']),
            )
        self.assertTrue(audio_path.exists())
        self.assertEqual(self.client.delete('/api/projects/' + project['id']).status_code, 204)
        self.assertFalse(audio_path.exists())

    def test_external_verification_blocks_narrative_and_unconfirmed_audio_blocks_full_mp3(self):
        project = self.client.post('/api/projects', json={'name':'对外内容','scenario':'legal_podcast','transform_mode':'enrich','verification_mode':'external_verify'}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title':'素材','content':'企业应当核验跨境数据传输、责任主体和备案要求。'}).json()
        document = self.client.get('/api/documents/' + source['id']).json()
        ids = [block['id'] for block in document['blocks'] if block['kind'] == 'paragraph']
        blocked = self.client.post(f"/api/projects/{project['id']}/skill-host-contents", json={'document_id':source['id'],'title':'播客稿','markdown':'# 播客稿\n\n正文。','source_block_ids':ids})
        self.assertEqual(blocked.status_code, 409)
        plan = self.client.post(f"/api/projects/{project['id']}/plans", json={'source_document_id':source['id']}).json()
        self.client.put(f"/api/plans/{plan['id']}/confirm", json={'chapters':plan['chapters']})
        state = self.client.get('/api/projects/' + project['id']).json()
        for task in state['tasks']:
            self.client.put('/api/tasks/' + task['id'], json={'status':'done'})
        content = self.client.post(f"/api/projects/{project['id']}/skill-host-contents", json={'document_id':source['id'],'title':'播客稿','markdown':'# 播客稿\n\n这是播客稿的正文内容，长度超过最小限制。','source_block_ids':ids}).json()
        audio = self.client.post(f"/api/projects/{project['id']}/audio-scripts", json={'narrative_content_id':content['id']}).json()
        self.assertEqual(self.client.post('/api/audio-outputs/' + audio['id'] + '/synthesize', json={}).status_code, 409)
        with patch.object(self.main, 'speech_chunk', return_value=(b'preview', 'test')):
            self.assertEqual(self.client.post('/api/audio-outputs/' + audio['id'] + '/synthesize', json={'preview':True}).status_code, 200)

    def test_daily_brief_subscription_collects_new_rss_items_once(self):
        xml = ET.fromstring('''<rss><channel><item><guid>one</guid><title>监管动态一</title><link>https://example.com/one</link><description>第一篇资讯正文，包含必要的法律规则说明。</description></item><item><guid>two</guid><title>监管动态二</title><link>https://example.com/two</link><description>第二篇资讯正文，包含必要的合规事项说明。</description></item></channel></rss>''')
        with patch.object(self.main.httpx, 'get') as get:
            response = type('Response', (), {'content': ET.tostring(xml), 'raise_for_status': lambda self: None})()
            get.return_value = response
            subscription = self.client.post('/api/daily-brief-subscriptions', json={'name':'监管资讯','feed_url':'https://example.com/feed.xml','daily_time':'08:00','max_items':2})
            self.assertEqual(subscription.status_code, 201)
            first = self.client.post('/api/daily-brief-subscriptions/' + subscription.json()['id'] + '/run')
            self.assertEqual(first.status_code, 200)
            self.assertEqual(first.json()['new_item_count'], 2)
            self.assertTrue(first.json()['project_id'])
            second = self.client.post('/api/daily-brief-subscriptions/' + subscription.json()['id'] + '/run')
            self.assertEqual(second.json()['new_item_count'], 0)

    def test_tts_payloads_and_invalid_success_response(self):
        import httpx
        config = {'tts_provider':'minimax', 'tts_api_key':'secret', 'tts_speed':0.9}
        response = httpx.Response(200, json={'base_resp':{'status_code':0}, 'data':{'audio':b'mp3'.hex()}}, request=httpx.Request('POST','https://example.com'))
        with patch.object(httpx, 'post', return_value=response) as post:
            audio, model = self.main.speech_chunk('测试', config)
            self.assertEqual(audio, b'mp3')
            self.assertEqual(model, 'speech-2.8-hd')
            self.assertEqual(post.call_args.kwargs['json']['voice_setting']['speed'], 0.9)
            self.assertTrue(post.call_args.args[0].endswith('/t2a_v2'))
        with patch.object(httpx, 'post', return_value=response):
            with self.assertRaises(self.main.HTTPException):
                self.main.speech_chunk('test', {'tts_base_url':'https://example.com','tts_api_key':'secret'})

    def test_macos_say_tts_uses_local_voice_without_api_key(self):
        with patch.object(self.main.sys, 'platform', 'darwin'), patch.object(self.main.shutil, 'which', return_value='/usr/bin/say'), patch.object(self.main.subprocess, 'run') as run:
            def make_output(args, **_kwargs):
                if '?' in args:
                    return subprocess.CompletedProcess(args, 0, stdout='Tingting            zh_CN    # 你好\n')
                target = Path(args[args.index('-o') + 1] if '-o' in args else args[-1])
                target.write_bytes(b'mp3')
                return subprocess.CompletedProcess(args, 0)
            run.side_effect = make_output
            audio, provider = self.main.speech_chunk('本地语音测试。', {'tts_provider':'macos_say','tts_voice':'Tingting','tts_speed':1})
        self.assertEqual(audio, b'mp3')
        self.assertEqual(provider, 'macos-say:Tingting')

    def test_edge_tts_uses_configured_voice_without_api_key(self):
        class FakeCommunicate:
            def __init__(self, text, voice, rate):
                self.voice = voice
                self.rate = rate

            async def save(self, path):
                Path(path).write_bytes(b'edge-mp3')

        fake_module = types.SimpleNamespace(Communicate=FakeCommunicate)
        with patch.object(self.main, 'edge_tts', fake_module):
            audio, provider = self.main.speech_chunk(
                '在线语音测试。',
                {'tts_provider': 'edge_tts', 'tts_voice': 'zh-CN-XiaoxiaoNeural', 'tts_speed': 1},
            )
        self.assertEqual(audio, b'edge-mp3')
        self.assertEqual(provider, 'edge-tts:zh-CN-XiaoxiaoNeural')

    def test_tts_settings_preview_does_not_require_text_model(self):
        with patch.object(self.main, 'get_internal_provider_settings', return_value={'tts_provider': 'edge_tts'}), patch.object(self.main, 'speech_chunk', return_value=(b'preview-mp3', 'edge-tts:test')):
            response = self.client.post('/api/settings/tts/test')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b'preview-mp3')
        self.assertEqual(response.headers['content-type'], 'audio/mpeg')

    def test_merge_ocr_lines_joins_sentences_and_drops_noise(self):
        lines = [
            '2026年9月25日，美国参议员宣布提出一项采购限制法案。',
            'McCormick办公室公布的名称',
            '是《保障国家安全系统免受中国光收发器影响法案》。',
            '微信扫一扫',
            '喜欢作者',
            '这项编号 S.5548 的法案',
        ]
        merged = self.main._merge_ocr_lines(lines)
        self.assertEqual(merged, [
            '2026年9月25日，美国参议员宣布提出一项采购限制法案。',
            'McCormick办公室公布的名称是《保障国家安全系统免受中国光收发器影响法案》。',
            '这项编号 S.5548 的法案',
        ])

    def test_provider_error_message_recognizes_authentication_failure(self):
        detail = '{"error":{"message":"Authentication Fails, Your api key: ****ZXQA is invalid (request_id: 11176610)","type":"authentication_error","code":"invalid_request_error"}}'
        message = self.main.provider_error_message(detail)
        self.assertIn('API Key 未通过服务商验证', message)
        self.assertIn('ZXQA', message)
        self.assertIn('空格', message)
        # 不再透传原始 JSON
        self.assertNotIn('request_id', message)

    def test_save_provider_strips_secret_whitespace(self):
        response = self.client.put('/api/settings/provider', json={
            'provider_preset': 'deepseek',
            'api_key': '  sk-secret-key  \n',
            'model_name': '',
            'allow_source_upload': True,
        })
        self.assertEqual(response.status_code, 200)
        stored = self.main.get_internal_provider_settings()
        self.assertEqual(stored['api_key'], 'sk-secret-key')

    def test_save_provider_keeps_keys_per_preset(self):
        base = {'model_name': '', 'allow_source_upload': True}
        # 1. 先给 deepseek 配一个 Key
        response = self.client.put('/api/settings/provider', json={**base, 'provider_preset': 'deepseek', 'api_key': 'sk-ds-123'})
        self.assertEqual(response.status_code, 200)
        # 2. 切到 openai：前端会把已配置的 Key 回显成占位符，这里填新 Key
        response = self.client.put('/api/settings/provider', json={**base, 'provider_preset': 'openai', 'api_key': 'sk-oai-456'})
        self.assertEqual(response.status_code, 200)
        stored = self.main.get_internal_provider_settings()
        self.assertEqual(stored['api_key'], 'sk-oai-456')
        self.assertEqual(stored['api_keys']['deepseek'], 'sk-ds-123')
        # 3. 切回 deepseek：占位符意味着沿用该服务商已存的 Key
        response = self.client.put('/api/settings/provider', json={**base, 'provider_preset': 'deepseek', 'api_key': '已配置（本地不回显）'})
        self.assertEqual(response.status_code, 200)
        stored = self.main.get_internal_provider_settings()
        self.assertEqual(stored['api_key'], 'sk-ds-123')
        # 4. GET 接口返回按服务商打码的 Key 表
        response = self.client.get('/api/settings/provider')
        data = response.json()
        self.assertEqual(data['api_keys'], {'deepseek': '已配置（本地不回显）', 'openai': '已配置（本地不回显）'})
        self.assertEqual(data['api_key'], '已配置（本地不回显）')

    def test_parse_pdf_uses_vision_ocr_for_image_pdf(self):
        try:
            import Vision  # noqa: F401
        except ImportError:
            self.skipTest('需要 macOS Vision 框架')
        from PIL import Image, ImageDraw, ImageFont
        img = Image.new('RGB', (1240, 900), 'white')
        draw = ImageDraw.Draw(img)
        font = None
        for candidate in ('/System/Library/Fonts/Supplemental/Arial.ttf', '/System/Library/Fonts/Helvetica.ttc'):
            try:
                font = ImageFont.truetype(candidate, 56)
                break
            except Exception:
                continue
        if font is None:
            self.skipTest('未找到可用的系统字体')
        draw.text((100, 150), 'Quantum Harbor Protocol Notes', fill='black', font=font)
        draw.text((100, 320), 'Section one describes the import pipeline.', fill='black', font=font)
        draw.text((100, 490), 'Section two lists fallback behaviour.', fill='black', font=font)
        path = self.main.DATA_DIR / 'ocr-fixture.pdf'
        img.save(path, 'PDF', resolution=100.0)
        try:
            blocks, _ = self.main.parse_pdf(path)
        finally:
            path.unlink(missing_ok=True)
        joined = ' '.join(block['text'] for block in blocks)
        self.assertIn('Quantum', joined)
        self.assertIn('fallback', joined)

    def test_fetch_article_filters_noise_and_keeps_headings(self):
        import httpx
        html = ('<html><head><meta property="og:title" content="当中国AI企业走上安理会讲台"/></head><body>'
                '<div id="js_content">'
                '<h2>为什么重要</h2>'
                '<p>中国代表在安理会会议上就人工智能治理阐明了立场，强调发展中国家应当拥有平等参与规则制定的权利。</p>'
                '<p>预览时标签不可点</p><p>不喜欢</p><p>微信扫一扫</p><p>关注该公众号</p><p>知道了</p><p>允许</p><p>使用完整服务</p><p>，</p><p>×</p>'
                '<p>文章还提到，主要国家正在竞争制定全球人工智能安全标准，技术出口管制成为博弈工具。</p>'
                '</div></body></html>')
        response = httpx.Response(200, headers={'content-type': 'text/html; charset=utf-8'}, text=html,
                                  request=httpx.Request('GET', 'https://mp.weixin.qq.com/s/xyz'))
        with patch.object(self.main.httpx, 'get', return_value=response):
            title, text = self.main.fetch_article('https://mp.weixin.qq.com/s/xyz')
        self.assertEqual(title, '当中国AI企业走上安理会讲台')
        for noise in ('预览时标签不可点', '不喜欢', '微信扫一扫', '关注该公众号', '知道了', '允许', '使用完整服务', '×'):
            self.assertNotIn(noise, text)
        self.assertNotIn('\n，\n', text)
        self.assertIn('## 为什么重要', text)
        self.assertIn('中国代表在安理会会议上', text)

    def test_web_page_title_extracts_og_title(self):
        import httpx
        project = self.client.post('/api/projects', json={'name': '标题抓取', 'verification_mode': 'source_only'}).json()
        page = httpx.Response(
            200,
            headers={'content-type': 'text/html; charset=utf-8'},
            text='<html><head><meta property="og:title" content="这篇讲 AI 治理的长文"/><title>微信 fallback 标题</title></head><body>正文</body></html>',
            request=httpx.Request('GET', 'https://mp.weixin.qq.com/s/abc'),
        )
        with patch.object(self.main.httpx, 'get', return_value=page):
            result = self.client.get('/api/web-page-title', params={'url': 'https://mp.weixin.qq.com/s/abc'})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['title'], '这篇讲 AI 治理的长文')
        # 无 og:title 时回落到 <title>
        page2 = httpx.Response(
            200,
            headers={'content-type': 'text/html; charset=utf-8'},
            text='<html><head><title>  只有 title 的页面 </title></head><body></body></html>',
            request=httpx.Request('GET', 'https://example.com/b'),
        )
        with patch.object(self.main.httpx, 'get', return_value=page2):
            result2 = self.client.get('/api/web-page-title', params={'url': 'https://example.com/b'})
        self.assertEqual(result2.json()['title'], '只有 title 的页面')
        # 非公开地址直接拒绝
        blocked = self.client.get('/api/web-page-title', params={'url': 'http://127.0.0.1:8080/'})
        self.assertEqual(blocked.status_code, 400)

    def test_quick_link_preview_does_not_create_project_before_confirmation(self):
        url = 'https://mp.weixin.qq.com/s/example'
        article = '这是公众号文章的正文。' * 30
        before = len(self.client.get('/api/projects').json())
        with patch.object(self.main, 'fetch_article', return_value=('测试文章', article)):
            preview = self.client.get('/api/link-preview', params={'url': url})
            self.assertEqual(preview.status_code, 200)
            self.assertEqual(preview.json()['title'], '测试文章')
            self.assertGreater(preview.json()['character_count'], 200)
            self.assertEqual(preview.json()['full_text'], article)
            self.assertEqual(preview.json()['paragraph_count'], 1)
            self.assertEqual(preview.json()['ending_excerpt'], article[-180:])
            self.assertEqual(len(self.client.get('/api/projects').json()), before)
            created = self.client.post('/api/projects/from-link', json={'url': url, 'scenario': 'daily_brief'})
        self.assertEqual(created.status_code, 201)
        detail = self.client.get('/api/projects/' + created.json()['project_id']).json()
        self.assertEqual(detail['project']['name'], '测试文章')
        self.assertEqual(detail['project']['scenario'], 'daily_brief')
        self.assertEqual(detail['documents'][0]['source_url'], url)
        shelf_item = next(item for item in self.client.get('/api/projects').json() if item['id'] == created.json()['project_id'])
        self.assertEqual(shelf_item['document_count'], 1)
        self.assertFalse(shelf_item['production_stage'].get('has_confirmed_outline', False))
        self.assertEqual(shelf_item['learning_progress'], {'done': 0, 'total': 0})

    def test_quick_link_rejects_short_or_private_page_without_empty_project(self):
        before = len(self.client.get('/api/projects').json())
        with patch.object(self.main, 'fetch_article', return_value=('验证页面', '请登录')):
            short = self.client.post('/api/projects/from-link', json={'url': 'https://example.com/paywall'})
        self.assertEqual(short.status_code, 422)
        private = self.client.get('/api/link-preview', params={'url': 'http://127.0.0.1:8080/private'})
        self.assertEqual(private.status_code, 400)
        self.assertEqual(len(self.client.get('/api/projects').json()), before)

    def test_web_search_requires_explicit_import_and_preserves_source(self):
        import httpx
        project = self.client.post('/api/projects', json={'name': '人工智能治理', 'scenario': 'topic_learning'}).json()
        search_response = httpx.Response(
            200,
            content='''<rss><channel><item><title>人工智能治理公开监管动态</title><link>https://example.com/news</link><description>&lt;b&gt;人工智能治理公开摘要&lt;/b&gt;</description></item></channel></rss>'''.encode('utf-8'),
            request=httpx.Request('GET', 'https://www.bing.com/search'),
        )
        with patch.object(self.main.httpx, 'get', return_value=search_response), \
             patch.object(self.main.httpx, 'post', side_effect=httpx.ConnectError('offline')):
            search = self.client.post('/api/web-search', json={'query': '人工智能治理'})
        self.assertEqual(search.status_code, 200)
        self.assertEqual(len(search.json()['results']), 1)
        self.assertEqual(self.client.get('/api/projects/' + project['id']).json()['documents'], [])

        page_response = httpx.Response(
            200,
            headers={'content-type': 'text/html; charset=utf-8'},
            text='<html><head><script>ignored()</script></head><body><h1>公开监管动态</h1><p>' + ('正文内容。' * 40) + '</p></body></html>',
            request=httpx.Request('GET', 'https://example.com/news'),
        )
        with patch.object(self.main.httpx, 'get', return_value=page_response):
            imported = self.client.post('/api/projects/' + project['id'] + '/web-sources', json={'title': '公开监管动态', 'url': 'https://example.com/news'})
        self.assertEqual(imported.status_code, 201)
        document = self.client.get('/api/documents/' + imported.json()['id']).json()
        self.assertEqual(document['source_url'], 'https://example.com/news')
        self.assertNotIn('ignored()', ' '.join(item['text'] for item in document['blocks']))
        blocked = self.client.post('/api/projects/' + project['id'] + '/web-sources', json={'title': '本机', 'url': 'http://127.0.0.1:8080/'})
        self.assertEqual(blocked.status_code, 400)

    def test_project_accepts_auto_research_mode(self):
        project = self.client.post('/api/projects', json={'name': '自动补充', 'web_research_mode': 'auto'})
        self.assertEqual(project.status_code, 201)
        data = self.client.get('/api/projects/' + project.json()['id']).json()['project']
        self.assertEqual(data['web_research_mode'], 'auto')

    def test_auto_research_imports_sources_and_reports_unusable_candidates(self):
        import httpx
        project = self.client.post('/api/projects', json={'name': '人工智能治理', 'scenario': 'topic_learning'}).json()
        search_response = httpx.Response(
            200,
            content='''<rss><channel><item><title>人工智能治理公开监管动态</title><link>https://example.com/news</link><description>人工智能治理公开摘要</description></item></channel></rss>'''.encode('utf-8'),
            request=httpx.Request('GET', 'https://www.bing.com/search'),
        )
        page_response = httpx.Response(
            200,
            headers={'content-type': 'text/html; charset=utf-8'},
            text='<html><body><h1>公开监管动态</h1><p>' + ('正文内容。' * 40) + '</p></body></html>',
            request=httpx.Request('GET', 'https://example.com/news'),
        )
        # 检索走 Bing RSS（DDG 的 post 被 mock 成离线），随后抓取正文入库。
        with patch.object(self.main.httpx, 'post', side_effect=httpx.ConnectError('offline')), \
             patch.object(self.main.httpx, 'get', side_effect=[search_response, page_response]):
            result = self.client.post('/api/projects/' + project['id'] + '/auto-research', json={'query': '人工智能治理'})
        self.assertEqual(result.status_code, 201)
        body = result.json()
        self.assertEqual(len(body['imported']), 1)
        self.assertEqual(body['failures'], [])
        self.assertIn('未经人工审阅', body['message'])
        # 自动导入的素材必须登记来源链接，否则无法追溯、也无法作为补充资料使用。
        document = self.client.get('/api/documents/' + body['imported'][0]['id']).json()
        self.assertEqual(document['source_url'], 'https://example.com/news')

        # 候选来源都提不出正文时应明确报错，而不是留下一份空素材。
        unusable = httpx.Response(
            200,
            headers={'content-type': 'text/html; charset=utf-8'},
            text='<html><body><p>需要登录后查看</p></body></html>',
            request=httpx.Request('GET', 'https://example.com/news'),
        )
        with patch.object(self.main.httpx, 'post', side_effect=httpx.ConnectError('offline')), \
             patch.object(self.main.httpx, 'get', side_effect=[search_response, unusable]):
            failed = self.client.post('/api/projects/' + project['id'] + '/auto-research', json={'query': '人工智能治理'})
        self.assertEqual(failed.status_code, 502)
        self.assertIn('手动粘贴', failed.json()['detail'])

    def test_real_ffmpeg_merge_and_measured_duration(self):
        path = self.main.DATA_DIR / 'fixture.mp3'
        subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','sine=frequency=440:duration=1','-y',str(path)], check=True)
        script = ('这是测试句子。\n' * 400).strip()
        chunks = self.main.split_speech_text(script)
        self.assertEqual(''.join(chunks), script)
        self.assertTrue(all(len(chunk) <= 1200 for chunk in chunks))
        with patch.object(self.main, 'speech_chunk', return_value=(path.read_bytes(), 'fixture')):
            output, _, duration = self.main.synthesize_audio(script, 'test', {})
        self.assertTrue(output.is_file())
        self.assertGreaterEqual(duration, len(chunks))
        self.assertLessEqual(duration, len(chunks) + 1)

    def test_macos_voice_detection_endpoint_and_resolution(self):
        fake = "Tingting               zh_CN    # 你好\nTingting (Enhanced)    zh_CN    # 你好\nMeijia                 zh_TW    # 您好\n"
        with patch.object(self.main, 'list_macos_zh_voices', return_value=self.main.__dict__ and [
            {'name': 'Tingting', 'lang': 'zh_cn', 'enhanced': False},
            {'name': 'Tingting (Enhanced)', 'lang': 'zh_cn', 'enhanced': True},
            {'name': 'Meijia', 'lang': 'zh_tw', 'enhanced': False},
        ]):
            response = self.client.get('/api/settings/macos-voices')
            self.assertEqual(response.status_code, 200)
            data = response.json()
            self.assertTrue(data['available'])
            self.assertTrue(data['enhanced_available'])
            self.assertEqual(len(data['voices']), 3)
            self.assertIn('增强版', data['hint'])
            # 基础音色自动升级为增强版
            self.assertEqual(self.main.resolve_macos_voice('Tingting'), 'Tingting (Enhanced)')
            # 未知音色回落到普通话（优先增强版）
            self.assertEqual(self.main.resolve_macos_voice('Nonexistent'), 'Tingting (Enhanced)')
            # 显式增强版保持不变
            self.assertEqual(self.main.resolve_macos_voice('Tingting (Enhanced)'), 'Tingting (Enhanced)')
        with patch.object(self.main, 'list_macos_zh_voices', return_value=[]):
            empty = self.client.get('/api/settings/macos-voices')
            self.assertFalse(empty.json()['available'])
            self.assertIn('系统设置', empty.json()['hint'])

    def test_demo_project_is_generic_and_legacy_cleanup(self):
        created = self.client.post('/api/demo-project')
        self.assertEqual(created.status_code, 201)
        first = created.json()
        self.assertTrue(first['created'])
        projects = self.client.get('/api/projects').json()
        names = [item['name'] for item in projects]
        self.assertIn(self.main.DEMO_PROJECT_NAME, names)
        self.assertNotIn(self.main.LEGACY_DEMO_PROJECT_NAME, names)
        detail = self.client.get('/api/projects/' + first['project_id']).json()
        script_texts = [item['script'] for item in detail['audio_outputs']]
        joined = ' '.join(script_texts)
        for term in ('法律', '律师', '备案', '合规'):
            self.assertNotIn(term, joined)
        self.assertIn('收藏', joined)
        # 重复调用幂等
        again = self.client.post('/api/demo-project').json()
        self.assertFalse(again['created'])
        self.assertEqual(again['project_id'], first['project_id'])

    def test_edge_tts_falls_back_to_macos_say_when_offline(self):
        class BrokenCommunicate:
            def __init__(self, *args, **kwargs):
                raise OSError('network unreachable')
        with patch.object(self.main.sys, 'platform', 'darwin'), patch.object(self.main.shutil, 'which', return_value='/usr/bin/say'), \
             patch.object(self.main.edge_tts, 'Communicate', BrokenCommunicate), patch.object(self.main.subprocess, 'run') as run:
            def make_output(args, **_kwargs):
                if '?' in args:
                    return subprocess.CompletedProcess(args, 0, stdout='Tingting            zh_CN    # 你好\n')
                target = Path(args[args.index('-o') + 1] if '-o' in args else args[-1])
                target.write_bytes(b'mp3')
                return subprocess.CompletedProcess(args, 0)
            run.side_effect = make_output
            audio, provider = self.main.speech_chunk('离线回落测试。', {'tts_provider':'edge_tts','tts_voice':'zh-CN-XiaoxiaoNeural','tts_speed':1})
        self.assertEqual(audio, b'mp3')
        self.assertEqual(provider, 'macos-say:Tingting')

    def test_merge_audiobook_builds_chaptered_m4b(self):
        with tempfile.TemporaryDirectory() as temp:
            paths = []
            for freq in (440, 660):
                chapter_path = Path(temp) / f'chapter-{freq}.mp3'
                subprocess.run(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', f'sine=frequency={freq}:duration=1', '-y', str(chapter_path)], check=True)
                paths.append(chapter_path)
            output = Path(temp) / 'book.m4b'
            result = self.main.merge_audiobook([('开端', paths[0]), ('收束', paths[1])], None, '测试有声书', output)
            self.assertTrue(result.is_file())
            probe = subprocess.run(['ffprobe', '-v', 'error', '-show_chapters', '-of', 'json', str(output)], check=True, capture_output=True, text=True)
            chapters = json.loads(probe.stdout)['chapters']
            self.assertEqual(len(chapters), 2)
            self.assertEqual(chapters[0]['tags']['title'], '第1章 开端')
            self.assertEqual(chapters[1]['tags']['title'], '第2章 收束')
            self.assertAlmostEqual(float(chapters[1]['start_time']), 1.0, delta=0.2)

    def test_audiobook_endpoint_exports_sectioned_m4b(self):
        fixture = self.main.DATA_DIR / 'fixture.mp3'
        subprocess.run(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=1', '-y', str(fixture)], check=True)
        project = self.client.post('/api/projects', json={'name': '有声书项目', 'verification_mode': 'source_only'}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title': '素材', 'content': '这是第一段材料内容。这是第二段材料内容。'}).json()
        document = self.client.get(f"/api/documents/{source['id']}").json()
        ids = [b['id'] for b in document['blocks'] if b['kind'] == 'paragraph']
        markdown = '# 讲稿\n\n## 第一节 开端\n\n开端内容，语调平缓。\n\n## 第二节 收束\n\n收束内容，给出行动建议。'
        content = self.client.post(f"/api/projects/{project['id']}/skill-host-contents", json={'document_id': source['id'], 'title': '章节讲稿', 'markdown': markdown, 'source_block_ids': ids}).json()
        confirmed = self.client.put(f"/api/narrative-contents/{content['id']}", json={'markdown': markdown, 'status': 'confirmed'})
        self.assertEqual(confirmed.status_code, 200)
        with tempfile.TemporaryDirectory() as temp:
            with patch.object(self.main, 'speech_chunk', return_value=(fixture.read_bytes(), 'fixture')):
                with patch.object(self.main, 'get_configured_export_directory', return_value=Path(temp)):
                    response = self.client.post(f"/api/projects/{project['id']}/audiobook")
            self.assertEqual(response.status_code, 201)
            data = response.json()
            self.assertEqual(data['chapters'], 2)
            audiobook = Path(data['audiobook_path'])
            self.assertEqual(audiobook.suffix, '.m4b')
            self.assertTrue(audiobook.is_file())
            probe = subprocess.run(['ffprobe', '-v', 'error', '-show_chapters', '-of', 'json', str(audiobook)], check=True, capture_output=True, text=True)
            chapters = json.loads(probe.stdout)['chapters']
            self.assertEqual(len(chapters), 2)
            self.assertIn('开端', chapters[0]['tags']['title'])
            # 容器品牌必须是 M4B，iPhone 图书 App 才会接收为有声书
            with open(response.json()['audiobook_path'], 'rb') as handle:
                head = handle.read(12)
                self.assertEqual(head[4:8], b'ftyp')
                self.assertEqual(head[8:12], b'M4B ')

    def test_direct_audiobook_endpoint_splits_by_source_headings(self):
        fixture = self.main.DATA_DIR / 'fixture.mp3'
        subprocess.run(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=1', '-y', str(fixture)], check=True)
        project = self.client.post('/api/projects', json={'name': '直读有声书', 'verification_mode': 'source_only'}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={
            'title': '直读素材',
            'content': '## 第一节 背景\n\n' + '背景段落，交代来龙去脉，铺陈事件的时间线与关键人物。' * 3 + '\n\n## 第二节 结论\n\n' + '结论段落，给出具体建议，并说明理由与适用范围。' * 3,
        }).json()
        direct = self.client.post(f"/api/projects/{project['id']}/documents/{source['id']}/direct-audio", json={})
        self.assertEqual(direct.status_code, 200)
        with tempfile.TemporaryDirectory() as temp:
            def fake_synthesize(script, title, settings):
                copy = Path(temp) / f'chapter-{abs(hash(title))}.mp3'
                copy.write_bytes(fixture.read_bytes())
                return copy, 'fixture', 1
            with patch.object(self.main, 'synthesize_audio', side_effect=fake_synthesize), \
                 patch.object(self.main, 'get_configured_export_directory', return_value=Path(temp)):
                response = self.client.post(f"/api/projects/{project['id']}/audiobook/direct")
            self.assertEqual(response.status_code, 201)
            data = response.json()
            self.assertEqual(data['chapters'], 2)
            audiobook = Path(data['audiobook_path'])
            self.assertTrue(audiobook.name.endswith('（原文朗读）.m4b'))
            probe = subprocess.run(['ffprobe', '-v', 'error', '-show_chapters', '-of', 'json', str(audiobook)], check=True, capture_output=True, text=True)
            chapters = json.loads(probe.stdout)['chapters']
            self.assertEqual(len(chapters), 2)
            self.assertIn('第一节 背景', chapters[0]['tags']['title'])
            self.assertIn('第二节 结论', chapters[1]['tags']['title'])
        # 逐章合成的中间文件应已清理，只留下 m4b
        leftovers = [p.name for p in Path(temp).glob('chapter-*.mp3')]
        self.assertEqual(leftovers, [])

    def test_direct_audiobook_chapters_merges_preamble_into_first_heading(self):
        blocks = [
            {'id': 'b0', 'kind': 'paragraph', 'text': '导语内容。', 'sequence_no': 1, 'heading_path': '', 'heading_level': 0, 'source_locator': ''},
            {'id': 'b1', 'kind': 'heading', 'text': '第一章', 'sequence_no': 2, 'heading_path': '', 'heading_level': 2, 'source_locator': ''},
            {'id': 'b2', 'kind': 'paragraph', 'text': '第一章正文。', 'sequence_no': 3, 'heading_path': '', 'heading_level': 0, 'source_locator': ''},
            {'id': 'b3', 'kind': 'heading', 'text': '第二章', 'sequence_no': 4, 'heading_path': '', 'heading_level': 2, 'source_locator': ''},
            {'id': 'b4', 'kind': 'paragraph', 'text': '第二章正文。', 'sequence_no': 5, 'heading_path': '', 'heading_level': 0, 'source_locator': ''},
        ]
        with patch.object(self.main, 'read_blocks', return_value=blocks):
            chapters = self.main.direct_audiobook_chapters('doc-1', '直读素材')
        self.assertEqual([title for title, _ in chapters], ['直读素材', '第一章', '第二章'])
        self.assertTrue(chapters[0][1].startswith('导语内容。'))
        self.assertTrue(chapters[1][1].startswith('第一章'))

    def test_material_map_extracts_generic_topic_signals(self):
        text = '收藏夹里的长文章越积越多。长文章需要整块时间，而碎片时间只能读短内容。' \
               '其实读完长文章的关键是换个方式遇见它：把长文章拆成小节，用碎片时间逐节消化。' \
               '长文章不该是负担，碎片时间也能读完长文章。'
        signals = self.main.extract_topic_signals(text)
        self.assertTrue(signals)
        names = [item['name'] for item in signals]
        self.assertTrue(any('长文章' in name or '文章' in name for name in names))
        blocks = [{'id': 'b1', 'kind': 'heading', 'heading_level': 1, 'text': '标题', 'sequence_no': 1, 'heading_path': '', 'source_locator': '段落 1'},
                  {'id': 'b2', 'kind': 'paragraph', 'text': text, 'sequence_no': 2, 'heading_path': '', 'source_locator': '段落 2'}]
        material_map = self.main.build_material_map(blocks)
        self.assertTrue(material_map['topics'])

    def test_audio_script_has_no_legal_outro(self):
        script = self.main.build_audio_script('# 标题\n\n正文第一段。', '测试讲稿', 'topic_learning', 5)
        self.assertNotIn('以上内容仅按已确认材料整理', script)
        self.assertIn('正文第一段。', script)

    def test_audio_script_section_intro_skips_document_title(self):
        title = '法律合规信息简讯（月刊）_2026年第9期_HW 2026-09-20 [Sent]'
        script = self.main.build_audio_script('## 第一章\n\n正文第一段。', title, 'legal_podcast', 5, '巴西能源转型四项法令')
        self.assertIn('接下来这一节，我们讲巴西能源转型四项法令。', script)
        # 章节音频不再念整篇文档标题（下划线、日期杠、[Sent] 状态标签）
        self.assertNotIn('简讯', script)
        self.assertNotIn('_', script)
        self.assertNotIn('[Sent]', script)

    def test_spoken_text_converts_machine_notation_to_speech(self):
        self.assertEqual(self.main.spoken_text('HW 2026-09-20 [Sent]'), 'HW 2026年9月20日')
        self.assertEqual(self.main.spoken_text('2026-09 发布'), '2026年9月 发布')
        self.assertEqual(self.main.spoken_text('文件_名'), '文件 名')
        self.assertIn('到', self.main.spoken_text('标准 50156-2021'))
        self.assertNotIn('·', self.main.spoken_text('甲·乙'))

    def test_whole_script_intro_reads_clean_title(self):
        script = self.main.build_audio_script('# 标题\n\n正文第一段。', '法律合规信息简讯（月刊）_2026年第9期_HW 2026-09-20 [Sent]', 'legal_podcast', 5)
        self.assertIn('2026年9月20日', script)
        self.assertNotIn('_', script)
        self.assertNotIn('[Sent]', script)

    def test_clean_pasted_markdown_strips_ai_chatter(self):
        pasted = '```markdown\n好的，以下是为你整理的讲稿：\n\n# 标题\n\n正文第一段。\n\n如果需要调整，请告诉我。\n```'
        cleaned = self.main.clean_pasted_markdown(pasted)
        self.assertEqual(cleaned, '# 标题\n\n正文第一段。')
        # 不误伤正文：以“当然”开头的普通句子要保留
        kept = self.main.clean_pasted_markdown('当然，也有例外情况。\n\n正文继续。')
        self.assertIn('当然，也有例外情况。', kept)

    def test_parse_ddg_results_and_unwrap(self):
        html = ('<a rel="nofollow" class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.com%2Fa">'
                '美国国会<b>光模块</b>法案</a>'
                '<a class="result__snippet" href="#">法案要点摘要</a>'
                '<a rel="nofollow" class="result__a" href="javascript:void(0)">坏链接</a>')
        results = self.main._parse_ddg_results(html, 5)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['url'], 'https://example.com/a')
        self.assertIn('光模块', results[0]['title'])
        self.assertEqual(results[0]['summary'], '法案要点摘要')

    def test_html_to_article_text_keeps_headings(self):
        html = '<html><script>var x=1;</script><body><h1>光模块法案</h1><p>第一段正文，足够长的一段话。</p><h2>背景</h2><p>第二段正文，说明法案出台的背景。</p><br/></body></html>'
        text = self.main.html_to_article_text(html)
        self.assertIn('## 光模块法案', text)
        self.assertIn('## 背景', text)
        self.assertIn('第一段正文', text)
        self.assertNotIn('var x=1', text)
        # 直接验证 parse_text 能识别出标题块
        import tempfile
        with tempfile.NamedTemporaryFile('w', suffix='.md', delete=False, encoding='utf-8') as handle:
            handle.write('标题\n\n' + text + '\n')
            path = Path(handle.name)
        blocks, _ = self.main.parse_text(path)
        headings = [b for b in blocks if b['kind'] == 'heading']
        self.assertGreaterEqual(len(headings), 2)

    def test_html_to_article_text_filters_ui_noise(self):
        html = ('<body><h1>法案解读</h1>'
                '<p>这是第一段正文内容，长度足够不会被过滤。</p>'
                '<p>预览时标签不可点</p>'
                '<p>修改于</p>'
                '<p>不喜欢</p>'
                '<p>相关推荐</p>'
                '<p>这是一句很短的评论。</p>'
                '<p>打开App阅读更多精彩内容</p>'
                '<p>这是结尾一段正文，同样长度足够保留下来。</p></body>')
        text = self.main.html_to_article_text(html)
        self.assertNotIn('预览时标签不可点', text)
        self.assertNotIn('修改于', text)
        self.assertNotIn('不喜欢', text)
        self.assertNotIn('相关推荐', text)
        self.assertNotIn('打开App', text)
        self.assertIn('这是第一段正文内容', text)
        self.assertIn('这是一句很短的评论。', text)
        self.assertIn('这是结尾一段正文', text)
        self.assertIn('## 法案解读', text)

    def test_delete_document_blocks_when_referenced(self):
        project = self.client.post('/api/projects', json={'name': '删除素材项目', 'verification_mode': 'source_only'}).json()
        project_id = project['id']
        source = self.client.post(f"/api/projects/{project_id}/text-sources", json={'title': '引用素材', 'content': '这是一段足够长的正文内容，用来通过最小长度校验。'}).json()
        document = self.client.get(f"/api/documents/{source['id']}").json()
        ids = [b['id'] for b in document['blocks'] if b['kind'] == 'paragraph']
        markdown = '# 讲稿\n\n## 第一节 开端\n\n引用素材生成的讲稿内容。'
        self.client.post(f"/api/projects/{project_id}/skill-host-contents", json={'document_id': source['id'], 'title': '章节讲稿', 'markdown': markdown, 'source_block_ids': ids})
        # 已有讲稿引用该素材，删除应被拦截
        blocked = self.client.delete(f"/api/projects/{project_id}/documents/{source['id']}")
        self.assertEqual(blocked.status_code, 409)
        # 未被引用的素材可以直接删除
        other = self.client.post(f"/api/projects/{project_id}/text-sources", json={'title': '待删素材', 'content': '这是另一段足够长的正文内容，用来通过最小长度校验。'}).json()
        deleted = self.client.delete(f"/api/projects/{project_id}/documents/{other['id']}")
        self.assertEqual(deleted.status_code, 200)
        missing = self.client.get(f"/api/documents/{other['id']}")
        self.assertEqual(missing.status_code, 404)

    def test_direct_audio_skips_rewrite_and_synthesizes(self):
        import httpx
        project = self.client.post('/api/projects', json={'name': '原文直读项目', 'verification_mode': 'source_only'}).json()
        source = self.client.post(f"/api/projects/{project['id']}/text-sources", json={'title': '公众号好文', 'content': '这是第一段原文内容，写得已经很好，不需要改写。' * 3 + '这是第二段原文内容，保持原样直接朗读即可。' * 3}).json()
        created = self.client.post(f"/api/projects/{project['id']}/documents/{source['id']}/direct-audio")
        self.assertEqual(created.status_code, 200)
        data = created.json()
        self.assertEqual(data['status'], 'script_ready')
        shelf = self.client.get('/api/projects').json()
        listed = next(item for item in shelf if item['id'] == project['id'])
        self.assertTrue(listed['production_stage']['has_direct_audio_script'])
        # 不经过任何改写：脚本与原文逐字一致
        output = self.client.get(f"/api/audio-outputs/{data['id']}").json()
        self.assertIn('不需要改写', output['script'])
        doc = self.client.get(f"/api/documents/{source['id']}").json()
        original = '\n\n'.join(b['text'] for b in doc['blocks'] if b['text'].strip())
        self.assertEqual(output['script'], original)
        # 无讲稿来源的直读输出不受"讲稿已确认"门禁限制，可直接合成 MP3
        fixture = self.main.DATA_DIR / 'fixture.mp3'
        if not fixture.is_file():
            subprocess.run(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=1', '-y', str(fixture)], check=True)
        with patch.object(self.main, 'speech_chunk', return_value=(fixture.read_bytes(), 'fixture')):
            synthesized = self.client.post(f"/api/audio-outputs/{data['id']}/synthesize", json={})
        self.assertEqual(synthesized.status_code, 200)
        self.assertEqual(synthesized.json()['status'], 'ready')
        # 幂等：再次点击复用同一条音频脚本，不产生重复卡片
        again = self.client.post(f"/api/projects/{project['id']}/documents/{source['id']}/direct-audio").json()
        self.assertEqual(again['id'], data['id'])
        self.assertTrue(again['reused'])

    def test_reveal_endpoint_rejects_outside_paths(self):
        response = self.client.post('/api/reveal-in-finder', json={'path': '/etc/hosts'})
        self.assertEqual(response.status_code, 403)

    def test_bundled_ffmpeg_registered_into_path(self):
        """随包 FFmpeg 目录必须注入 PATH：否则「生成 MP3」在没装 FFmpeg 的 Windows 上必然 400。"""
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / 'tools' / 'ffmpeg'
            exe_dir = root / 'ffmpeg-9.0.2-essentials_build' / 'bin'
            exe_dir.mkdir(parents=True)
            exe_name = 'ffmpeg.exe' if sys.platform == 'win32' else 'ffmpeg'
            (exe_dir / exe_name).write_bytes(b'fake')
            # POSIX 上 shutil.which 要求可执行位；Windows 靠扩展名判断，此行在 win 上是 no-op
            os.chmod(exe_dir / exe_name, 0o755)
            original = os.environ.get('PATH', '')
            try:
                found = self.main.register_bundled_ffmpeg([root])
                self.assertEqual([exe_dir], found)
                self.assertEqual(os.environ['PATH'].split(os.pathsep)[0], str(exe_dir))
                self.assertTrue(shutil.which('ffmpeg'))
                # Windows 下 which 返回的实际大小写可能是 .EXE，按路径比较不区分大小写
                self.assertEqual(Path(shutil.which('ffmpeg')).resolve(), (exe_dir / exe_name).resolve())
                # 本机 PATHEXT 被改成只剩 .CPL 时也必须兜底补上 .EXE，否则 which 永远找不到 ffmpeg.exe（仅 Windows 有 PATHEXT）
                if sys.platform == 'win32':
                    self.assertIn('.EXE', [item.upper() for item in os.environ.get('PATHEXT', '').split(os.pathsep)])
            finally:
                os.environ['PATH'] = original

    def test_bundled_ffmpeg_missing_root_is_noop(self):
        original = os.environ.get('PATH', '')
        try:
            found = self.main.register_bundled_ffmpeg([Path('Z:/definitely-not-exist') / 'ffmpeg'])
            self.assertEqual(found, [])
            self.assertEqual(os.environ.get('PATH', ''), original)
        finally:
            os.environ['PATH'] = original
